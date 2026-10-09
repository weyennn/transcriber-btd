"""Pembentuk dokumen: naskah menjadi PDF dan DOCX bergaya naskah dinas.

PDF memakai templat Typst pada templates/naskah-dinas.typ. Templat itu tidak
mencetak keterangan alat dan tidak mencetak penanda akhir dokumen, berbeda dari
pembentuk PDF notulen yang mencetak "Dibuat oleh Mesin Transcribe · perlu
verifikasi manual" pada setiap halaman. DOCX dibentuk dengan susunan yang sama
supaya naskah yang disunting tetap sebentuk dengan yang dibagikan.

Seluruh teks naskah dikirim ke Typst sebagai string, bukan sebagai markup,
sehingga tanda seperti $, _, *, atau kurung siku di dalam naskah tidak pernah
ditafsirkan sebagai perintah tata letak.

Pembentuk PDF dijalankan sebagai proses terpisah terhadap perkakas `typst` yang
sudah ada di dalam image. Bila perkakas itu tidak tersedia, dokumen tidak
dibentuk dan pemanggil menerima penjelasan, bukan berkas yang rusak.
"""

from __future__ import annotations

import re
import shutil
import subprocess
from pathlib import Path

from src.pandora_documents.composer import RENCANA, Naskah

TEMPLATE_PATH = Path(__file__).with_name("templates") / "naskah-dinas.typ"


class PerkakasTidakTersedia(RuntimeError):
    """Perkakas pembentuk PDF tidak ada di lingkungan ini."""


# ── Naskah Markdown menjadi blok ─────────────────────────────────────────
def _polos(teks: str) -> str:
    """Buang penanda penekanan Markdown; naskah dinas tidak menebalkan istilah."""
    teks = re.sub(r"\*\*(.+?)\*\*", r"\1", teks)
    teks = re.sub(r"(?<![\w*])\*([^*\n]+?)\*(?![\w*])", r"\1", teks)
    return teks.replace("`", "").strip()


def urai(markdown: str) -> list[tuple]:
    """Urai subset Markdown isi naskah menjadi blok.

    Blok: ("judul", tingkat, teks), ("paragraf", teks), ("butir", [teks]),
    ("nomor", [teks]), ("tabel", [[sel]]).
    """
    blok: list[tuple] = []
    baris = markdown.replace("\r\n", "\n").split("\n")
    paragraf: list[str] = []

    def tutup_paragraf():
        if paragraf:
            blok.append(("paragraf", _polos(" ".join(paragraf))))
            paragraf.clear()

    i = 0
    while i < len(baris):
        isi = baris[i].strip()
        if not isi or re.fullmatch(r"[-*_]{3,}", isi):
            tutup_paragraf()
            i += 1
            continue

        judul = re.match(r"^(#{1,6})\s+(.+)$", isi)
        if judul:
            tutup_paragraf()
            blok.append(("judul", len(judul.group(1)), _polos(judul.group(2))))
            i += 1
            continue

        if isi.startswith("|"):
            tutup_paragraf()
            rows: list[list[str]] = []
            while i < len(baris) and baris[i].strip().startswith("|"):
                sel = [_polos(c) for c in baris[i].strip().strip("|").split("|")]
                if sel and not all(re.fullmatch(r"[-: ]*", c) for c in sel):
                    rows.append(sel)
                i += 1
            if rows:
                kolom = max(len(r) for r in rows)
                blok.append(("tabel", [r + [""] * (kolom - len(r)) for r in rows]))
            continue

        for jenis, pola in (("butir", r"^[-*]\s+(.+)$"), ("nomor", r"^\d+[.)]\s+(.+)$")):
            if re.match(pola, isi):
                tutup_paragraf()
                butir: list[str] = []
                while i < len(baris) and (cocok := re.match(pola, baris[i].strip())):
                    butir.append(_polos(cocok.group(1)))
                    i += 1
                blok.append((jenis, butir))
                break
        else:
            paragraf.append(isi)
            i += 1
    tutup_paragraf()
    return blok


def _tempat_tanggal(kepala: dict | None) -> str:
    """Baris tempat dan tanggal di atas tanda tangan; kosong bila belum lengkap."""
    isian = kepala or {}
    tempat, tanggal = (str(isian.get(k) or "").strip() for k in ("tempat", "tanggal"))
    return f"{tempat}, {tanggal}" if tempat and tanggal else ""


# ── Typst ────────────────────────────────────────────────────────────────
def _str(teks: str) -> str:
    """Tulis teks sebagai string Typst."""
    aman = teks.replace("\\", "\\\\").replace('"', '\\"')
    return '"' + re.sub(r"\s+", " ", aman).strip() + '"'


def _tabel_typst(rows: list[list[str]]) -> list[str]:
    """Tabel bergaris buku: garis hanya di atas, di bawah kepala, dan di bawah."""
    kolom = len(rows[0])
    keluar = [
        f"#table(columns: {kolom},",
        "  table.hline(stroke: .8pt),",
        "  table.header(" + ", ".join(f"[#{_str(c)}]" for c in rows[0]) + "),",
        "  table.hline(stroke: .5pt),",
    ]
    for baris in rows[1:]:
        keluar.append("  " + ", ".join(f"[#{_str(c)}]" for c in baris) + ",")
    keluar.append("  table.hline(stroke: .8pt),")
    keluar.append(")")
    return keluar


def isi_ke_typst(markdown: str) -> str:
    """Konversi isi satu bagian naskah ke Typst."""
    keluar: list[str] = []
    for blok in urai(markdown):
        if blok[0] == "judul":
            keluar.append(f"=== #{_str(blok[2])}")
        elif blok[0] == "paragraf":
            keluar.append(f"#{_str(blok[1])}")
        elif blok[0] in ("butir", "nomor"):
            tanda = "-" if blok[0] == "butir" else "+"
            keluar.extend(f"{tanda} #{_str(butir)}" for butir in blok[1])
        elif blok[0] == "tabel":
            keluar.extend(_tabel_typst(blok[1]))
        keluar.append("")
    return "\n".join(keluar)


def sumber_typst(naskah: Naskah, kepala: dict | None = None) -> str:
    """Isi templat dengan naskah yang sudah dikonversi."""
    baris_kepala = ",\n".join(
        f"  ({_str(label)}, {_str(nilai)})" for label, nilai in naskah.kepala(kepala)
    )
    badan: list[str] = []
    for butir in naskah.bagian:
        badan.append(f"== #{_str(f'{butir.nomor}. {butir.judul}')}")
        badan.append("")
        badan.append(isi_ke_typst(butir.isi))
    return (
        TEMPLATE_PATH.read_text(encoding="utf-8")
        .replace("__JENIS__", _str(naskah.nama_jenis))
        .replace("__JUDUL__", _str(naskah.judul))
        .replace("__KEPALA__", baris_kepala + ("," if baris_kepala else ""))
        .replace("__CATATAN__", _str(naskah.catatan_sumber()))
        .replace("__TEMPAT_TANGGAL__", _str(_tempat_tanggal(kepala)))
        .replace("__PENANDA_TANGAN__", _str(RENCANA[naskah.jenis]["penanda_tangan"]))
        .replace("__BODY__", "\n".join(badan))
    )


def bentuk_pdf(naskah: Naskah, keluaran: str | Path, kepala: dict | None = None,
               timeout: int = 120) -> str:
    """Bentuk PDF dari naskah. Kembalikan jalur PDF."""
    typst = shutil.which("typst")
    if not typst:
        raise PerkakasTidakTersedia(
            "Perkakas typst tidak ada di lingkungan ini, PDF tidak dibentuk."
        )
    target = Path(keluaran)
    target.parent.mkdir(parents=True, exist_ok=True)
    sumber = target.with_suffix(".typ")
    sumber.write_text(sumber_typst(naskah, kepala), encoding="utf-8")
    try:
        subprocess.run(
            [typst, "compile", str(sumber), str(target)],
            check=True, capture_output=True, text=True, timeout=timeout,
        )
    except subprocess.CalledProcessError as exc:
        rincian = (exc.stderr or exc.stdout or "compile gagal").strip()[-1200:]
        raise RuntimeError(f"Typst gagal membentuk PDF: {rincian}") from exc
    finally:
        # Sumber Typst adalah berkas kerja, bukan keluaran untuk pengguna.
        sumber.unlink(missing_ok=True)
    return str(target)


# ── DOCX ─────────────────────────────────────────────────────────────────
FONT_DOCX = "Times New Roman"
TITIK_ISIAN = "." * 60


def bentuk_docx(naskah: Naskah, keluaran: str | Path, kepala: dict | None = None) -> str:
    """Bentuk DOCX dengan susunan yang sama seperti PDF. Kembalikan jalur DOCX."""
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt

    doc = Document()
    seksi = doc.sections[0]
    seksi.page_width, seksi.page_height = Cm(21), Cm(29.7)
    seksi.top_margin = seksi.bottom_margin = Cm(2.5)
    seksi.left_margin = seksi.right_margin = Cm(2.5)

    gaya = doc.styles["Normal"]
    gaya.font.name = FONT_DOCX
    gaya.font.size = Pt(12)
    gaya.element.rPr.rFonts.set(qn("w:eastAsia"), FONT_DOCX)

    def paragraf(teks="", rata=WD_ALIGN_PARAGRAPH.JUSTIFY, tebal=False, ukuran=None,
                 menjorok=False, sebelum=0, sesudah=6, gaya_par=None):
        par = doc.add_paragraph(style=gaya_par)
        par.alignment = rata
        fmt = par.paragraph_format
        fmt.space_before, fmt.space_after = Pt(sebelum), Pt(sesudah)
        fmt.line_spacing = 1.15
        if menjorok:
            fmt.first_line_indent = Cm(1.25)
        if teks:
            run = par.add_run(teks)
            run.bold = tebal
            run.font.name = FONT_DOCX
            if ukuran:
                run.font.size = Pt(ukuran)
        return par

    def garis_sel(sel, **sisi):
        batas = OxmlElement("w:tcBorders")
        for nama, tebal in sisi.items():
            garis = OxmlElement(f"w:{nama}")
            garis.set(qn("w:val"), "single")
            garis.set(qn("w:sz"), str(tebal))
            garis.set(qn("w:color"), "000000")
            batas.append(garis)
        sel._tc.get_or_add_tcPr().append(batas)

    def isi_sel(sel, teks, tebal=False, rata=WD_ALIGN_PARAGRAPH.LEFT):
        par = sel.paragraphs[0]
        par.alignment = rata
        par.paragraph_format.space_after = Pt(2)
        run = par.add_run(teks)
        run.bold = tebal
        run.font.name = FONT_DOCX

    tengah = WD_ALIGN_PARAGRAPH.CENTER
    paragraf(naskah.nama_jenis.upper(), rata=tengah, tebal=True, ukuran=14, sesudah=2)
    paragraf(naskah.judul, rata=tengah, tebal=True, ukuran=12, sesudah=12)

    baris_kepala = naskah.kepala(kepala)
    if baris_kepala:
        tabel = doc.add_table(rows=len(baris_kepala), cols=3)
        tabel.autofit = False
        for baris, (label, nilai) in zip(tabel.rows, baris_kepala):
            for sel, lebar in zip(baris.cells, (Cm(4.2), Cm(0.5), Cm(11.3))):
                sel.width = lebar
            isi_sel(baris.cells[0], label)
            isi_sel(baris.cells[1], ":")
            isi_sel(baris.cells[2], nilai or TITIK_ISIAN)
        for sel in tabel.rows[-1].cells:
            garis_sel(sel, bottom=8)
        paragraf(sesudah=4)

    for butir in naskah.bagian:
        par = paragraf(f"{butir.nomor}. {butir.judul}", rata=WD_ALIGN_PARAGRAPH.LEFT,
                       tebal=True, sebelum=10, sesudah=4)
        par.paragraph_format.keep_with_next = True
        for blok in urai(butir.isi):
            if blok[0] == "judul":
                paragraf(blok[2], rata=WD_ALIGN_PARAGRAPH.LEFT, tebal=True, sebelum=6, sesudah=3)
            elif blok[0] == "paragraf":
                paragraf(blok[1], menjorok=True)
            elif blok[0] in ("butir", "nomor"):
                gaya_par = "List Bullet" if blok[0] == "butir" else "List Number"
                for urutan, teks in enumerate(blok[1], 1):
                    if blok[0] == "nomor":
                        # Penomoran ditulis langsung supaya tiap senarai mulai dari 1;
                        # gaya List Number bawaan melanjutkan hitungan antarsenarai.
                        par = paragraf(f"{urutan}.\t{teks}", sesudah=3)
                        par.paragraph_format.left_indent = Cm(1.25)
                        par.paragraph_format.first_line_indent = Cm(-0.75)
                    else:
                        paragraf(teks, sesudah=3, gaya_par=gaya_par)
            elif blok[0] == "tabel":
                rows = blok[1]
                tabel = doc.add_table(rows=len(rows), cols=len(rows[0]))
                tabel.alignment = WD_TABLE_ALIGNMENT.CENTER
                for r, (baris, isi_baris) in enumerate(zip(tabel.rows, rows)):
                    for sel, teks in zip(baris.cells, isi_baris):
                        isi_sel(sel, teks, tebal=(r == 0))
                        if r == 0:
                            garis_sel(sel, top=8, bottom=4)
                        if r == len(rows) - 1:
                            garis_sel(sel, bottom=8)
                paragraf(sesudah=4)

    paragraf(naskah.catatan_sumber(), ukuran=10, sebelum=12, sesudah=12)

    tempat_tanggal = _tempat_tanggal(kepala) or "." * 34
    kiri = WD_ALIGN_PARAGRAPH.LEFT
    blok_ttd = [tempat_tanggal, f"{RENCANA[naskah.jenis]['penanda_tangan']},", "", "", "",
                "." * 34]
    for urutan, teks in enumerate(blok_ttd):
        par = paragraf(teks, rata=kiri, sesudah=0)
        par.paragraph_format.left_indent = Cm(9.5)
        par.paragraph_format.keep_together = True
        if urutan < len(blok_ttd) - 1:
            par.paragraph_format.keep_with_next = True

    footer = seksi.footer.paragraphs[0]
    footer.alignment = tengah
    run = footer.add_run()
    for jenis_medan, teks in (("begin", None), (None, "PAGE"), ("end", None)):
        if jenis_medan:
            medan = OxmlElement("w:fldChar")
            medan.set(qn("w:fldCharType"), jenis_medan)
        else:
            medan = OxmlElement("w:instrText")
            medan.set(qn("xml:space"), "preserve")
            medan.text = teks
        run._r.append(medan)
    run.font.name = FONT_DOCX
    run.font.size = Pt(10)

    target = Path(keluaran)
    target.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(target))
    return str(target)
