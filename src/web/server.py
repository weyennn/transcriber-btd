"""FastAPI server — dashboard transcribe via browser (IP local).

Run:
    .venv\\Scripts\\python.exe -m src.web.server            # dev, 127.0.0.1
    .venv\\Scripts\\python.exe -m src.web.server --host 0.0.0.0 --port 8765
"""

import argparse
import json
import os
import re
import socket
import threading
import time
import webbrowser
from pathlib import Path

from fastapi import Depends, FastAPI, File, HTTPException, Request, UploadFile
from fastapi.responses import FileResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from starlette.background import BackgroundTask

from src.core.engine import TranscribeEngine
from src.notulen.ai_generator import ai_config, check_ai_router
from src.utils.config import load_config, save_config
from src.utils.ffmpeg_checker import check_ffmpeg
from src.utils.paths import HASIL_DIR, list_history, result_dir_for, upload_dir_for
from src.web.auth import auth_me, callback as auth_callback, get_current_user, is_archive_admin, login as auth_login, logout as auth_logout, user_id
from src.web.job_manager import JobManager
from src.web import job_store

WEB_DIR = Path(__file__).resolve().parent
TEMPLATE_DIR = WEB_DIR / "templates"
STATIC_DIR = WEB_DIR / "static"
from src.utils.paths import UPLOAD_DIR  # noqa: E402  (menghormati DATA_DIR)

# Patch WDAC paling awal (R-02 kajian) — sebelum import faster_whisper
TranscribeEngine.apply_wdac_patch()

app = FastAPI(title="Pandora", version="0.1.0")
jobs = JobManager()
app.mount("/static", StaticFiles(directory=str(STATIC_DIR)), name="static")

ALLOWED_EXT = {
    ".mp3", ".wav", ".m4a", ".flac", ".aac", ".ogg",
    ".wma", ".aiff", ".aif", ".amr", ".opus",
    ".webm", ".mp4", ".m4v", ".mov", ".mkv", ".avi",
}
ALLOWED_MODELS = {"tiny", "base", "small", "medium"}
ALLOWED_LANGUAGES = {"id", "en"}
AI_ENABLED = os.environ.get("AI_ENABLED", "false").strip().lower() == "true"


def _env_int(name: str, default: int) -> int:
    value = os.environ.get(name, str(default)).strip()
    try:
        parsed = int(value)
    except ValueError:
        return default
    return parsed if parsed > 0 else default


MAX_UPLOAD_BYTES = _env_int("MAX_UPLOAD_BYTES", 2 * 1024**3)


def _within_directory(path: str, directory: Path) -> bool:
    try:
        return Path(path).resolve().is_relative_to(directory.resolve())
    except OSError:
        return False


def _validate_upload_path(path: str) -> str:
    resolved = str(Path(path).resolve())
    if not _within_directory(resolved, UPLOAD_DIR):
        raise HTTPException(400, "File audio harus berasal dari folder upload")
    if not os.path.isfile(resolved):
        raise HTTPException(404, f"File audio tidak ditemukan: {resolved}")
    return resolved


def _owner_id(user: dict) -> str:
    return "legacy" if user.get("dev") else user_id(user)


def _user_upload_dir(user: dict) -> Path:
    return upload_dir_for(_owner_id(user))


def _user_result_dir(user: dict) -> Path:
    return result_dir_for(_owner_id(user))


def _result_folder_for_user(value: str, owner_id: str) -> Path:
    if not re.fullmatch(r"\d{2}_[\w.\- ]+", value):
        raise HTTPException(400, "Nama folder hasil tidak valid")
    root = result_dir_for(owner_id).resolve()
    folder = (root / value).resolve()
    if not folder.is_relative_to(root):
        raise HTTPException(400, "Folder di luar direktori hasil tidak diizinkan")
    return folder


def _sanitize_filename(name: str) -> str:
    name = os.path.basename(name.replace("\\", "/"))
    name = re.sub(r"[^\w.\- ]", "_", name)
    return name.strip() or "audio"

def _resolve_result_folder(value: str) -> Path:
    """Resolve a result folder name or legacy absolute path safely."""
    candidate = Path(value)
    if not candidate.is_absolute():
        candidate = HASIL_DIR / candidate.name
    return Path(os.path.abspath(os.path.normpath(str(candidate))))


@app.get("/")
def index():
    return FileResponse(TEMPLATE_DIR / "index.html")


@app.get("/api/auth/me")
def auth_me_route(request: Request):
    return auth_me(request)


@app.get("/api/auth/login")
def auth_login_route(request: Request):
    return auth_login(request)


@app.get("/api/auth/callback")
def auth_callback_route(request: Request):
    return auth_callback(request)


@app.post("/api/auth/logout")
def auth_logout_route():
    return auth_logout()


# ── Env check (R-09) ─────────────────────────────────────────────────────
@app.get("/api/env")
def env(_user: dict = Depends(get_current_user)):
    ff = check_ffmpeg()
    model_cached = TranscribeEngine.is_model_cached("medium")
    ai = ai_config()
    ai["router"] = check_ai_router() if AI_ENABLED else {
        "reachable": False,
        "ok": False,
        "status_code": None,
        "error": "AI_ENABLED=false",
    }
    return {
        "ffmpeg": ff,
        "model_cached": model_cached,
        "model_default": "medium",
        "output_dir": str(_user_result_dir(_user)),
        "ai": ai,
        "ai_enabled": AI_ENABLED,
        "max_upload_bytes": MAX_UPLOAD_BYTES,
        "settings": load_config(_owner_id(_user)),
    }


@app.get("/healthz")
def healthz():
    return {"ok": True}


@app.post("/api/upload")
async def upload(file: UploadFile = File(...), _user: dict = Depends(get_current_user)):
    ext = os.path.splitext(file.filename or "")[1].lower()
    if ext not in ALLOWED_EXT:
        raise HTTPException(
            400,
            f"Ekstensi {ext or '(kosong)'} tidak didukung. "
            f"Izinkan: {', '.join(sorted(ALLOWED_EXT))}",
        )

    upload_dir = _user_upload_dir(_user)
    upload_dir.mkdir(parents=True, exist_ok=True)
    safe_name = _sanitize_filename(file.filename or "audio")

    nofollow = getattr(os, "O_NOFOLLOW", 0)
    stem_ = Path(safe_name).stem
    suffix = Path(safe_name).suffix
    for index in range(1, 1_000_000):
        candidate = safe_name if index == 1 else f"{stem_}-{index}{suffix}"
        dest = upload_dir / candidate
        try:
            fd = os.open(dest, os.O_WRONLY | os.O_CREAT | os.O_EXCL | nofollow, 0o600)
            break
        except FileExistsError:
            continue
    else:
        raise HTTPException(409, "Tidak dapat menyediakan nama file upload")

    size = 0
    try:
        with os.fdopen(fd, "wb") as out:
            while chunk := await file.read(1024 * 1024):
                size += len(chunk)
                if size > MAX_UPLOAD_BYTES:
                    raise HTTPException(413, "Ukuran file melebihi batas upload")
                out.write(chunk)
    except HTTPException:
        dest.unlink(missing_ok=True)
        raise

    return {"path": str(dest), "filename": dest.name, "size": size}


# ── Transcribe ───────────────────────────────────────────────────────────
@app.post("/api/transcribe")
def start_transcribe(body: dict, _user: dict = Depends(get_current_user)):
    audio_path = body.get("audio_path") or body.get("file_path")
    if not audio_path:
        raise HTTPException(400, "audio_path wajib diisi")

    model = body.get("model", "medium")
    language = body.get("language", "id")
    if model not in ALLOWED_MODELS:
        raise HTTPException(400, f"Model tidak valid. Pilih: {', '.join(sorted(ALLOWED_MODELS))}")
    if language not in ALLOWED_LANGUAGES:
        raise HTTPException(400, "Bahasa tidak valid. Pilih: id atau en")

    audio_path = _validate_upload_path(audio_path)
    if not Path(audio_path).resolve().is_relative_to(_user_upload_dir(_user).resolve()):
        raise HTTPException(404, "File audio tidak ditemukan")
    output_dir = _user_result_dir(_user)
    output_dir.mkdir(parents=True, exist_ok=True)
    config = {
        "model": model,
        "language": language,
        "device": "cpu",
        "compute_type": "int8",
        "with_md": False,
        "output_dir": TranscribeEngine.auto_folder(Path(audio_path).stem, output_dir),
    }

    try:
        save_config({"model": model, "language": language,
                     "device": config["device"], "compute_type": config["compute_type"]},
                    owner_id=_owner_id(_user))
    except OSError:
        pass  # gagal simpan bukan error fatal

    job = jobs.create(audio_path, config, owner_id=_owner_id(_user))
    return job.to_dict()



@app.get("/api/jobs/active")
def active_jobs(_user: dict = Depends(get_current_user)):
    """Job aktif untuk recovery setelah browser refresh."""
    return {"items": jobs.active(_owner_id(_user))}


@app.get("/api/jobs/queue-status")
def queue_status(_user: dict = Depends(get_current_user)):
    """Status antrean agregat tanpa membocorkan owner atau path job."""
    return job_store.queue_status()


@app.get("/api/jobs/{job_id}")
def job_status(job_id: str, _user: dict = Depends(get_current_user)):
    job = jobs.get(job_id, _owner_id(_user))
    if not job:
        raise HTTPException(404, "Job tidak ditemukan")
    payload = job.to_dict() if hasattr(job, "to_dict") else dict(job)
    position = job_store.queue_position(job_id)
    if position:
        payload["queue_ahead"] = position["ahead"]
        payload["queue_position"] = position["ahead"] + 1
    return payload


@app.get("/api/jobs/{job_id}/stream")
def job_stream(job_id: str, _user: dict = Depends(get_current_user)):
    if not jobs.get(job_id, _owner_id(_user)):
        raise HTTPException(404, "Job tidak ditemukan")
    return StreamingResponse(
        jobs.event_stream(job_id, _owner_id(_user)),
        media_type="text/event-stream",
        headers={
            "Cache-Control": "no-cache",
            "X-Accel-Buffering": "no",
        },
    )


@app.post("/api/jobs/{job_id}/cancel")
def job_cancel(job_id: str, _user: dict = Depends(get_current_user)):
    ok = jobs.cancel(job_id, _owner_id(_user))
    if not ok:
        raise HTTPException(404, "Job tidak ditemukan")
    return {"ok": True}




@app.post("/api/notulen/ai")
def notulen_ai(body: dict, _user: dict = Depends(get_current_user)):
    """Buat notulen AI dari transkrip di folder hasil → DOCX."""
    if not AI_ENABLED:
        raise HTTPException(503, "Fitur AI notulen dimatikan di server ini (AI_ENABLED=false)")
    output_dir = body.get("output_dir") or body.get("folder") or body.get("folder_name")
    if not output_dir:
        raise HTTPException(400, "output_dir wajib diisi")
    fpath = _result_folder_for_user(Path(output_dir).name, _owner_id(_user))
    resolved = fpath.resolve()
    if not fpath.is_dir():
        raise HTTPException(404, "Folder hasil tidak ditemukan")
    if not resolved.is_relative_to(_user_result_dir(_user).resolve()):
        raise HTTPException(404, "Folder hasil tidak ditemukan")
    return jobs.create_ai(str(fpath), owner_id=_owner_id(_user)).to_dict()

@app.post("/api/transcript/cleanup")
def transcript_cleanup(body: dict, _user: dict = Depends(get_current_user)):
    """Rapikan transkrip dengan AI tanpa mengubah file transkrip asli."""
    if not AI_ENABLED:
        raise HTTPException(
            503,
            "Fitur AI dimatikan di server ini (AI_ENABLED=false)",
        )
    output_dir = body.get("output_dir") or body.get("folder") or body.get("folder_name")
    if not output_dir:
        raise HTTPException(400, "output_dir wajib diisi")
    fpath = _result_folder_for_user(Path(output_dir).name, _owner_id(_user))
    resolved = fpath.resolve()
    if not fpath.is_dir():
        raise HTTPException(404, "Folder hasil tidak ditemukan")
    if not resolved.is_relative_to(_user_result_dir(_user).resolve()):
        raise HTTPException(404, "Folder hasil tidak ditemukan")
    return jobs.create_cleanup(str(fpath), owner_id=_owner_id(_user)).to_dict()


# ── History (hasil transkripsi) ──────────────────────────────────────────
@app.get("/api/history")
def history(_user: dict = Depends(get_current_user)):
    items = []
    result_root = _user_result_dir(_user)
    for folder in list_history(result_root):
        folder_path = Path(folder)
        resolved_folder = folder_path.resolve()
        if not resolved_folder.is_relative_to(result_root.resolve()):
            continue
        files = [f.name for f in resolved_folder.iterdir()
                 if f.is_file() and (
                     not f.name.lower().endswith(".md")
                     or f.name in {"notulen_ai.md", "transkrip_rapi_ai.md"}
                 )]
        files.sort(key=lambda n: (n.startswith("transkrip"), n))
        items.append({"name": resolved_folder.name, "files": files})
    return {"items": items}


@app.get("/api/admin/archive")
def legacy_archive(_user: dict = Depends(get_current_user)):
    if not is_archive_admin(_user):
        raise HTTPException(403, "Akses arsip lama hanya untuk admin")
    archive_root = (HASIL_DIR / "_archive_legacy").resolve()
    if not archive_root.is_relative_to(HASIL_DIR.resolve()) or not archive_root.is_dir():
        return {"items": []}
    return {"items": [{"name": item.name} for item in sorted(archive_root.iterdir()) if item.is_dir()]}


@app.get("/api/admin/archive/{folder_name}/file/{file_name}")
def legacy_archive_file(folder_name: str, file_name: str, download: bool = False,
                        _user: dict = Depends(get_current_user)):
    if not is_archive_admin(_user):
        raise HTTPException(403, "Akses arsip lama hanya untuk admin")
    if not re.fullmatch(r"\d{2}_[\w\- ]+", folder_name):
        raise HTTPException(404, "File tidak ditemukan")
    root = (HASIL_DIR / "_archive_legacy").resolve()
    folder = (root / folder_name).resolve()
    name = os.path.basename(file_name)
    path = (folder / name).resolve()
    if not folder.is_relative_to(root) or not path.is_relative_to(folder) or not path.is_file():
        raise HTTPException(404, "File tidak ditemukan")
    if download:
        return FileResponse(path, filename=path.name)
    try:
        content = path.read_text(encoding="utf-8")
    except UnicodeDecodeError:
        raise HTTPException(415, "Preview hanya mendukung file teks; gunakan unduh")
    return {"name": path.name, "content": content, "size": path.stat().st_size}

@app.delete("/api/history/{folder_name}")
def delete_history(folder_name: str, _user: dict = Depends(get_current_user)):
    """Hapus satu folder hasil yang dipilih dari history."""
    if not re.match(r"^\d{2}_[\w.\- ]+$", folder_name):
        raise HTTPException(400, "Nama folder tidak valid")
    folder = _result_folder_for_user(folder_name, _owner_id(_user))
    if not folder.is_dir():
        raise HTTPException(404, "Folder hasil tidak ditemukan")
    import shutil
    shutil.rmtree(folder)
    return {"ok": True, "deleted": folder_name}


@app.get("/api/history/{folder_name}/file/{file_name}")
def history_file(folder_name: str, file_name: str, download: bool = False, _user: dict = Depends(get_current_user)):
    if not re.match(r"^\d{2}_[\w.\- ]+$", folder_name):
        raise HTTPException(400, "Nama folder tidak valid")
    file_name = os.path.basename(file_name)
    folder = _result_folder_for_user(folder_name, _owner_id(_user))
    fpath = (folder / file_name).resolve()
    if not fpath.is_relative_to(_user_result_dir(_user).resolve()) or not fpath.is_relative_to(folder):
        raise HTTPException(404, "File tidak ditemukan")
    if not folder.is_dir():
        raise HTTPException(404, "Folder hasil tidak ditemukan")

    transcript_stem = re.sub(r"[^\w.\- ]", "", folder_name).strip()[:50] or "transkrip"
    legacy_transcript_stem = re.sub(r"[^\w\- ]", "", folder_name).strip()[:50] or "transkrip"
    transcript_docx_names = {
        f"Transkrip_{folder_name}.docx",
        f"Transkrip_{transcript_stem}.docx",
        f"Transkrip_{legacy_transcript_stem}.docx",
    }
    if download and file_name in transcript_docx_names and not fpath.exists():
        transcript_path = (folder / "transkrip.txt").resolve()
        if not transcript_path.is_relative_to(_user_result_dir(_user).resolve()) or not transcript_path.is_file():
            raise HTTPException(404, "Transkrip asli tidak ditemukan")
        try:
            from docx import Document
        except ImportError as exc:
            raise HTTPException(503, "Generator DOCX tidak tersedia") from exc

        document = Document()
        document.add_heading(f"Transkrip — {folder_name}", level=1)
        transcript_text = " ".join(
            line.strip() for line in transcript_path.read_text(encoding="utf-8").splitlines()
            if line.strip()
        )
        document.add_paragraph(transcript_text)
        import tempfile
        with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as temp_file:
            temp_path = Path(temp_file.name)
        try:
            document.save(temp_path)
            return FileResponse(
                temp_path,
                filename=f"Transkrip_{folder_name}.docx",
                media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                background=BackgroundTask(temp_path.unlink, missing_ok=True),
            )
        except Exception:
            temp_path.unlink(missing_ok=True)
            raise

    if download:
        media_type = (
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            if fpath.suffix.lower() == ".docx"
            else "application/pdf" if fpath.suffix.lower() == ".pdf"
            else "text/plain; charset=utf-8"
        )
        return FileResponse(fpath, filename=fpath.name, media_type=media_type)

    try:
        content = fpath.read_text(encoding="utf-8")
    except UnicodeDecodeError:
        content = fpath.read_text(encoding="utf-8", errors="replace")
    return {"name": file_name, "content": content, "size": fpath.stat().st_size}


# ── Launch helper ────────────────────────────────────────────────────────
def get_local_ips() -> list:
    ips = ["127.0.0.1"]
    try:
        s = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
        s.connect(("8.8.8.8", 80))
        ips.append(s.getsockname()[0])
        s.close()
    except OSError:
        pass
    return list(dict.fromkeys(ips))


def run_server(host="127.0.0.1", port=8765, open_browser=True):
    import uvicorn

    if host in ("0.0.0.0", "::"):
        urls = [f"http://{ip}:{port}" for ip in get_local_ips()]
    else:
        urls = [f"http://{host}:{port}"]

    print("=" * 56)
    print("  PANDORA — CONVERSATIONS WORTH KEEPING")
    print("=" * 56)
    for u in urls:
        print(f"  Buka di browser: {u}")
    print("  Tekan Ctrl+C untuk menghentikan server.")
    print("=" * 56)

    if open_browser:
        def _open():
            time.sleep(1.2)
            webbrowser.open(urls[0])
        threading.Thread(target=_open, daemon=True).start()

    uvicorn.run(app, host=host, port=port, log_level="warning")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Pandora — audio transcription workspace")
    parser.add_argument("--host", default="127.0.0.1",
                        help="Bind host (default 127.0.0.1; 0.0.0.0 = semua)")
    parser.add_argument("--port", type=int, default=8765)
    parser.add_argument("--no-browser", action="store_true",
                        help="Jangan buka browser otomatis")
    args = parser.parse_args()
    run_server(host=args.host, port=args.port, open_browser=not args.no_browser)
