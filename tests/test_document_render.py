"""Uji pembentuk dokumen resmi: blok naskah, sumber Typst, DOCX, dan PDF.

Uji PDF hanya berjalan bila perkakas typst tersedia (ada di dalam image).
"""

import shutil

import pytest
from docx import Document

from src.pandora_documents import render
from src.pandora_documents.composer import Bagian, Naskah

ISI = (
    "Rapat membahas anggaran $5 juta untuk unit_kerja [tidak jelas] dan #tagar @tim.\n\n"
    "1. Capaian pertama 94 persen.\n"
    "2. Capaian kedua 57 persen.\n\n"
    "- Butir \"berkutip\" dengan garis miring terbalik \\ di dalamnya.\n\n"
    "| Tindak lanjut | Penanggung jawab |\n|---|---|\n| Melengkapi data | belum ditetapkan |\n"
)
KEPALA = {"tanggal": "Jumat, 9 Oktober 2026", "tempat": "Yogyakarta", "unit": "", "tujuan": ""}


def naskah_contoh():
    return Naskah("Monitoring Capaian", "laporan_kegiatan",
                  [Bagian("I", "Dasar dan Maksud Pelaksanaan", "", ISI),
                   Bagian("II", "Pelaksanaan", "", "Kegiatan berjalan menurut rekaman.")])


def test_urai_mengenali_blok_naskah():
    blok = render.urai("### Subjudul\n\nBaris satu\nbaris **dua**.\n\n" + ISI)
    assert [b[0] for b in blok] == ["judul", "paragraf", "paragraf", "nomor", "butir", "tabel"]
    assert blok[1] == ("paragraf", "Baris satu baris dua.")
    assert blok[3][1] == ["Capaian pertama 94 persen.", "Capaian kedua 57 persen."]
    assert blok[5][1] == [["Tindak lanjut", "Penanggung jawab"], ["Melengkapi data", "belum ditetapkan"]]


def test_sumber_typst_tanpa_penanda_alat_dan_tanpa_markup_dari_naskah():
    sumber = render.sumber_typst(naskah_contoh(), KEPALA)
    assert "__" not in sumber.replace("unit_kerja", "")
    for terlarang in ("Dibuat oleh", "verifikasi manual", "Akhir Dokumen", "rgb("):
        assert terlarang not in sumber
    assert '== #"I. Dasar dan Maksud Pelaksanaan"' in sumber
    # Teks naskah masuk sebagai string, jadi tanda markup tidak perlu dan tidak boleh lolos mentah.
    assert '#"Rapat membahas anggaran $5 juta untuk unit_kerja [tidak jelas] dan #tagar @tim."' in sumber
    assert '\\"berkutip\\"' in sumber and "terbalik \\\\ di" in sumber
    assert '("Unit penyusun", "")' in sumber
    assert '#let tempat-tanggal = "Yogyakarta, Jumat, 9 Oktober 2026"' in sumber
    assert "table.hline(stroke: .8pt)" in sumber


def test_docx_bersusunan_naskah_dinas(tmp_path):
    jalur = render.bentuk_docx(naskah_contoh(), tmp_path / "naskah.docx", KEPALA)
    doc = Document(jalur)
    teks = [p.text for p in doc.paragraphs]
    assert teks[0] == "LAPORAN KEGIATAN" and teks[1] == "Monitoring Capaian"
    assert "I. Dasar dan Maksud Pelaksanaan" in teks and "II. Pelaksanaan" in teks
    assert "1.\tCapaian pertama 94 persen." in teks
    assert "Yogyakarta, Jumat, 9 Oktober 2026" in teks and "Penyusun," in teks
    semua = "\n".join(teks)
    for terlarang in ("Akhir Dokumen", "Dibuat oleh", "Daftar Isi", "—"):
        assert terlarang not in semua
    kepala = doc.tables[0]
    assert [sel.text for sel in kepala.rows[0].cells] == ["Hari, tanggal", ":", "Jumat, 9 Oktober 2026"]
    assert set(kepala.rows[2].cells[2].text) == {"."}  # isian kosong untuk diisi tangan
    seksi = doc.sections[0]
    assert round(seksi.page_width.cm, 1) == 21.0 and round(seksi.left_margin.cm, 1) == 2.5


def test_pdf_ditolak_dengan_penjelasan_bila_typst_tidak_ada(tmp_path, monkeypatch):
    monkeypatch.setattr(render.shutil, "which", lambda nama: None)
    with pytest.raises(render.PerkakasTidakTersedia):
        render.bentuk_pdf(naskah_contoh(), tmp_path / "naskah.pdf")
    assert not list(tmp_path.iterdir())


@pytest.mark.skipif(not shutil.which("typst"), reason="perkakas typst tidak tersedia")
def test_pdf_terbentuk_dari_naskah_bertanda_khusus(tmp_path):
    jalur = render.bentuk_pdf(naskah_contoh(), tmp_path / "naskah.pdf", KEPALA)
    assert (tmp_path / "naskah.pdf").read_bytes().startswith(b"%PDF")
    assert jalur == str(tmp_path / "naskah.pdf")
    assert not (tmp_path / "naskah.typ").exists()
